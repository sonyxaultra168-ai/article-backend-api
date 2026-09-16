import os
import re
import io
import base64
import uuid
import glob
import subprocess
import urllib.request
import stat
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import google.generativeai as genai
import requests
from bs4 import BeautifulSoup
import docx
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
import imageio_ffmpeg

app = Flask(__name__)
CORS(app)

CONFIG_DIR = '/tmp'
YTDLP_URL = "https://github.com/yt-dlp/yt-dlp/releases/latest/download/yt-dlp"
YTDLP_EXE = os.path.join(CONFIG_DIR, "yt-dlp")

@app.route('/')
def index():
    return "Article Rewriter API is Online!"

@app.route('/fetch_url', methods=['POST'])
def fetch_url():
    url = request.json.get('url', '').strip()
    if not url: return jsonify({'error': 'សូមបញ្ចូលលីងវេបសាយជាមុនសិន!'})
    try:
        res = requests.get(url, headers={'User-Agent': 'Mozilla/5.0'}, timeout=10)
        res.raise_for_status()
        soup = BeautifulSoup(res.text, 'html.parser')
        paragraphs = soup.find_all('p')
        text = '\n\n'.join([p.get_text().strip() for p in paragraphs if p.get_text().strip()])
        if not text: text = soup.get_text(separator='\n', strip=True)
        return jsonify({'success': True, 'text': text[:20000]}) 
    except Exception as e:
        return jsonify({'error': f"មិនអាចទាញទិន្នន័យពីវេបសាយនេះបានទេ! ⚠️"})

@app.route('/update_ytdlp', methods=['POST'])
def update_ytdlp():
    try:
        urllib.request.urlretrieve(YTDLP_URL, YTDLP_EXE)
        st = os.stat(YTDLP_EXE)
        os.chmod(YTDLP_EXE, st.st_mode | stat.S_IEXEC)
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': f"មិនអាច Update បានទេ: {str(e)}"})

@app.route('/download_audio', methods=['POST'])
def download_audio():
    url = request.json.get('url', '').strip()
    if not url: return jsonify({'error': 'សូមបញ្ចូល Link ជាមុនសិន!'})
    if not os.path.exists(YTDLP_EXE):
        return jsonify({'error': 'សូមចុចប៊ូតុង Update (ពណ៌ខៀវ) ដើម្បីទាញយកកម្មវិធី YT-DLP ជាមុនសិន!'})

    temp_dir = os.path.join(CONFIG_DIR, 'temp')
    os.makedirs(temp_dir, exist_ok=True)
    out_template = os.path.join(temp_dir, f"{uuid.uuid4().hex}.%(ext)s")

    # ទាញយកផ្លូវរបស់ ffmpeg ចេញពី Library ដែលយើងទើបបន្ថែម
    ffmpeg_path = imageio_ffmpeg.get_ffmpeg_exe()

    # បន្ថែម --ffmpeg-location ដើម្បីអោយ yt-dlp ស្គាល់ម៉ាស៊ីនបំប្លែង
    cmd = [YTDLP_EXE, '--ffmpeg-location', ffmpeg_path, '-x', '--audio-format', 'mp3', '--no-playlist', '-o', out_template, url]
    
    try:
        subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        files = glob.glob(out_template.replace('%(ext)s', 'mp3'))
        if not files: return jsonify({'error': 'មិនអាចទាញយកសំឡេងពីប្រភពនេះបានទេ! (អាចជាប់ Privacy)'})
        
        with open(files[0], 'rb') as f:
            encoded = base64.b64encode(f.read()).decode('utf-8')
            
        os.remove(files[0]) 
        
        return jsonify({'success': True, 'audio_base64': 'data:audio/mp3;base64,' + encoded})
    except Exception as e:
        return jsonify({'error': 'មានបញ្ហាក្នុងការទាញយក! សូមប្រាកដថា Link ត្រូវនិងមិនមែនជា Private វីដេអូ។'})

@app.route('/scan_models', methods=['POST'])
def scan_models():
    api_key = request.form.get('api_key', '').strip()
    if not api_key: return jsonify({'error': 'សូមបញ្ចូល API Key ជាមុនសិន!'})
    try:
        genai.configure(api_key=api_key)
        models = []
        for m in genai.list_models():
            if 'generateContent' in m.supported_generation_methods:
                name = m.name.lower()
                if 'gemini' in name and 'vision' not in name:
                    display = m.display_name
                    if 'exp' in name or 'preview' in name: display = f"[សាកល្បង] {display}"
                    models.append({'id': m.name, 'name': display})
        
        if not models:
            return jsonify({'error': 'រកមិនឃើញម៉ូឌែល Gemini ដែលត្រឹមត្រូវទេ! សូមពិនិត្យ Key ម្តងទៀត។'})
            
        models.sort(key=lambda x: '0' if '2.0' in x['name'] else '1' if '1.5' in x['name'] else '2')
        return jsonify({'success': True, 'models': models})
    except Exception as e:
        return jsonify({'error': "API Key ខុស! សូមពិនិត្យមើលម្តងទៀត។ ❌"})

@app.route('/rewrite_article', methods=['POST'])
def rewrite_article():
    data = request.json
    api_key = data.get('api_key', '').strip()
    model_id = data.get('model_id', '').strip()
    original_text = data.get('text', '').strip()
    mode = data.get('mode', 'dynamic')
    audio_base64 = data.get('audio_base64')
    audio_mime = data.get('audio_mime', 'audio/mp3')
    country = data.get('country', 'កម្ពុជា (Cambodia)').strip()
    tone = data.get('tone', 'default').strip()

    if not api_key: return jsonify({'error': 'សូមកំណត់ API Key របស់ Gemini ជាមុនសិន!'})
    if not model_id: return jsonify({'error': 'សូមជ្រើសរើសម៉ាស៊ីន Gemini ជាមុនសិន!'})
    if not original_text and not audio_base64: return jsonify({'error': 'សូមបញ្ចូលទិន្នន័យ (អត្ថបទ ឬ MP3) ជាមុនសិន!'})

    tone_instruction = "Write in an elite, professional, and highly engaging journalistic tone."
    if tone == 'formal': tone_instruction = "Write in a highly formal, objective, and authoritative tone suitable for official government or political news."
    elif tone == 'entertainment': tone_instruction = "Write in a catchy, sensational, and entertaining gossip style, keeping readers hooked with dramatic flair."
    elif tone == 'emotional': tone_instruction = "Write in a deeply emotional, empathetic, and storytelling tone that touches the reader's heart."
    elif tone == 'investigative': tone_instruction = "Write in an investigative, mysterious, and analytical tone, focusing on uncovering hidden truths and building suspense."

    base_rules = "You are an elite, award-winning professional journalist and a master storyteller writing exclusively in the **Khmer language**.\nCRITICAL RULES:\n1. PERFECT KHMER: Use flawless standard Khmer spelling and grammar.\n2. SHATTER & INVERT THE STRUCTURE (CRITICAL): You are STRICTLY FORBIDDEN from keeping the original paragraph order. YOU MUST PARSE ALL FACTS FIRST, THEN REBUILD:\n   - RULE A: NEVER start your article with the same information as the original text.\n   - RULE B: INVERTED PYRAMID: You MUST extract the most shocking statistics, the core demand, or the final conclusion from the BOTTOM/MIDDLE of the original text and FORCE it to be your FIRST paragraph.\n   - RULE C: Group remaining facts logically. DO NOT paraphrase sentence-by-sentence.\n"

    if country != 'កម្ពុជា (Cambodia)':
        if mode == 'generate_new':
            base_rules += f"\n3. MULTILINGUAL NATIVE RESEARCH WORKFLOW (Target: '{country}'):\n   - STEP 1: Research this topic internally using the native language of '{country}'.\n   - STEP 2: Process those facts into English.\n   - STEP 3: Translate to standard Khmer.\n   - STEP 4: COMPLETELY SHATTER any dry encyclopedic tone."
        else:
            base_rules += f"\n3. STRICT 3-STEP MULTILINGUAL WORKFLOW (Source: '{country}'):\n   - Translate the foreign input into English, then to standard Khmer.\n   - COMPLETELY SHATTER IT. DO NOT output a direct translation."
    else:
        if mode == 'generate_new':
            base_rules += "\n3. LOCAL CAMBODIAN RESEARCH WORKFLOW:\n   - Research the topic deeply using local context before writing."

    enhancement_rules = f"\n   - TONE & STYLE: {tone_instruction}\n   - ABSOLUTE RESTRUCTURING MANDATE (CRITICAL): You are STRICTLY FORBIDDEN from following the paragraph order of the original text. You MUST extract the final conclusion, core demand, or most critical data from the MIDDLE or END of the source text and force it to be your opening paragraph. Completely shuffle and re-weave the remaining facts into a brand new sequence.\n   - SOCIAL MEDIA PROMO: Generate a short, highly engaging Facebook/Telegram caption summarising the article, along with 3-5 trending hashtags in Khmer.\n   - IMAGE PROMPTS: Generate 2 or 3 hyper-realistic AI image generation prompts in ENGLISH (e.g., for Midjourney) that perfectly visually represent the climax or main subject of the story."

    title_rule = "\n6. MYSTERY TITLES & DYNAMIC SCORING: Generate exactly 5 Khmer titles. Use the 'Curiosity Gap' technique (hide the main secret so they MUST click). Do NOT output fake hardcoded scores. Genuinely evaluate each title's clickability. Give the most mysterious title the highest score (e.g. [9.9/10]) and score the rest based on actual quality. Sort highest to lowest."

    structure_rules = ""
    if mode in ['custom_paragraphs', '4_paragraphs']:
        para_count = data.get('para_count', 4)
        structure_rules = f"\nSTORYTELLING & FORMAT RULES:\n4. Masterful Narrative Structure:\n   - DESTROY THE ORIGINAL SEQUENCE: Gather all facts and present them in a completely new narrative arc.\n   - STRICT PARAGRAPH COUNT: EXACTLY {para_count} paragraphs.{enhancement_rules}{title_rule}"
    elif mode == 'dynamic':
        structure_rules = f"\nSTORYTELLING & FORMAT RULES:\n4. Dynamic Narrative Structure:\n   - ADAPT TO LENGTH: Cover every original fact but in a NEW logical flow.\n   - DESTROY THE ORIGINAL SEQUENCE: Gather all facts and present them in a completely new narrative arc.{enhancement_rules}{title_rule}"
    elif mode == 'narration':
        structure_rules = f"\nSTORYTELLING & FORMAT RULES:\n4. Direct Narration Extraction:\n   - Write a comprehensive news article recounting their exact story.\n   - EXTRACT & REBUILD: Find the climax or the most shocking claim of their story and put it FIRST. Do not invent facts.\n   - FIX ERRORS: Fix all speech-to-text spelling/grammar mistakes to standard Khmer.{enhancement_rules}{title_rule}"
    elif mode == 'quote':
        attr_phrase = data.get('attribution_phrase', '')
        person_name = data.get('person_name', '')
        para_count = data.get('paragraph_count', 'auto')
        context_text = data.get('context', '').strip()
        post_date = data.get('post_date', '').strip() 
        n_str = f"EXACTLY {int(para_count)+1} paragraphs." if str(para_count).isdigit() and int(para_count)>0 else "Expand naturally."
        ctx_str = f"\n   - CRITICAL CONTEXT: '{context_text}'." if context_text else ""
        date_str = f"\n   - DATE CONTEXT: '{post_date}'." if post_date else ""
        structure_rules = f"\nSTORYTELLING & FORMAT RULES:\n4. Quote Expansion:\n   - {n_str}{ctx_str}{date_str}\n   - MUST introduce the quote EXACTLY as: \"{attr_phrase} {person_name} បានរៀបរាប់យ៉ាងដូច្នេះថា៖ [INSERT QUOTE HERE]\".\n   - Expand on its meaning using the inverted pyramid method.{enhancement_rules}{title_rule}"
    elif mode == 'generate_new':
        para_count = data.get('paragraph_count', 'auto')
        n_str = f"EXACTLY {int(para_count)} paragraphs." if str(para_count).isdigit() and int(para_count)>0 else "Expand naturally."
        structure_rules = f"\nSTORYTELLING & FORMAT RULES:\n4. Original Article Generation:\n   - BRAND NEW article from scratch based on topic.\n   - {n_str}\n   - SHATTER DRY FACTS: Weave facts into a magnetic story.{enhancement_rules}{title_rule}"

    output_format = "\n\nOUTPUT FORMAT (CRITICAL):\nYou MUST respond EXACTLY using the tags below. Do NOT use JSON. Just output plain text with these tags:\n\n[ARTICLE]\nYour flawless Khmer article goes here.\n[/ARTICLE]\n\n[SOCIAL]\nCaption and #hashtags go here.\n[/SOCIAL]\n\n[IMAGE_PROMPTS]\n1. English Prompt 1...\n2. English Prompt 2...\n[/IMAGE_PROMPTS]\n\n[TITLES]\n1. [Score/10] Title here...\n2. [Score/10] Title here...\n3. [Score/10] Title here...\n4. [Score/10] Title here...\n5. [Score/10] Title here...\n[/TITLES]"

    system_prompt = base_rules + structure_rules + output_format

    try:
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel(model_id, system_instruction=system_prompt)
        
        contents = []
        if audio_base64:
            b64_data = audio_base64.split(',')[1] if ',' in audio_base64 else audio_base64
            audio_bytes = base64.b64decode(b64_data)
            contents.append({"mime_type": audio_mime, "data": audio_bytes})
        
        if original_text:
            contents.append(original_text)
        
        if not contents:
            return jsonify({'error': 'គ្មានទិន្នន័យដើម្បីវិភាគទេ!'})

        response = model.generate_content(
            contents,
            generation_config=genai.GenerationConfig(temperature=0.8, top_p=0.9, top_k=40)
        )
        
        result_text = response.text.strip()
        article_match = re.search(r'\[ARTICLE\](.*?)\[/ARTICLE\]', result_text, re.DOTALL | re.IGNORECASE)
        social_match = re.search(r'\[SOCIAL\](.*?)\[/SOCIAL\]', result_text, re.DOTALL | re.IGNORECASE)
        image_match = re.search(r'\[IMAGE_PROMPTS\](.*?)\[/IMAGE_PROMPTS\]', result_text, re.DOTALL | re.IGNORECASE)
        titles_match = re.search(r'\[TITLES\](.*?)\[/TITLES\]', result_text, re.DOTALL | re.IGNORECASE)

        if not article_match or not titles_match:
            return jsonify({'error': "ម៉ាស៊ីន AI សរសេរខុសទម្រង់! សូមចុចសរសេរម្ដងទៀត។"})

        rewritten_text = article_match.group(1).strip()
        social_text = social_match.group(1).strip() if social_match else ""
        image_text = image_match.group(1).strip() if image_match else ""

        if social_text:
            rewritten_text += f"\n\n--- 📝 សម្រាប់ Social Media (Caption & Hashtags) ---\n{social_text}"
        if image_text:
            rewritten_text += f"\n\n--- 🎨 គំនិតរូបភាព (AI Image Prompts) ---\n{image_text}"

        titles_raw = titles_match.group(1).strip().split('\n')
        catchy_titles = []
        for t in titles_raw:
            t = t.strip()
            if t:
                clean_t = re.sub(r'^(\d+\.|\-|\*)\s*', '', t).strip()
                clean_t = re.sub(r'^["\']|["\']$', '', clean_t).strip()
                if clean_t: catchy_titles.append(clean_t)
        
        return jsonify({
            'success': True,
            'rewritten_text': rewritten_text,
            'catchy_titles': catchy_titles[:5]
        })

    except Exception as e:
        err_msg = str(e).lower()
        if 'quota' in err_msg or '429' in err_msg:
            return jsonify({'error': "API Key អស់កូតាហើយ! សូមប្តូរ Key ថ្មី។ ⚠️"})
        return jsonify({'error': f"បរាជ័យ: {str(e)[:100]}"})

@app.route('/download_word', methods=['POST'])
def download_word():
    try:
        import docx
        from docx.shared import Pt, RGBColor
        from docx.oxml.ns import qn
        
        data = request.json
        text = data.get('text', '')
        titles = data.get('titles', [])

        doc = docx.Document()
        
        def add_khmer_run(paragraph, p_text, size_pt, is_bold=False, color_rgb=None):
            run = paragraph.add_run(p_text)
            run.font.name = 'Khmer OS Siemreap'
            run.font.size = Pt(size_pt)
            if is_bold: run.font.bold = True
            if color_rgb: run.font.color.rgb = color_rgb
            run._element.rPr.rFonts.set(qn('w:ascii'), 'Khmer OS Siemreap')
            run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Khmer OS Siemreap')
            run._element.rPr.rFonts.set(qn('w:cs'), 'Khmer OS Siemreap')
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Khmer OS Siemreap')

        if titles:
            h1 = doc.add_heading(level=1)
            add_khmer_run(h1, 'ចំណងជើង (Titles)', 14, is_bold=True, color_rgb=RGBColor(0, 0, 0))
            for t in titles:
                try:
                    p = doc.add_paragraph(style='List Bullet')
                    add_khmer_run(p, t, 12)
                except:
                    p = doc.add_paragraph()
                    add_khmer_run(p, f"- {t}", 12)
                    
        h2 = doc.add_heading(level=1)
        add_khmer_run(h2, 'អត្ថបទ (Article & Prompts)', 14, is_bold=True, color_rgb=RGBColor(0, 0, 0))
        for para in text.split('\n'):
            p_text = para.strip()
            if p_text:
                p = doc.add_paragraph()
                add_khmer_run(p, p_text, 12)

        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        return send_file(
            file_stream, 
            as_attachment=True, 
            download_name='Article_Rewriter_PRO.docx', 
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

    except Exception as e:
        return jsonify({'error': f"កំហុសប្រព័ន្ធ: {str(e)}"}), 500

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 10000))
    app.run(host='0.0.0.0', port=port)